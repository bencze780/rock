from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def create_presentation():
    # Prezentáció létrehozása
    prs = Presentation()

    # --- 1. DIA: Címlap ---
    slide_layout = prs.slide_layouts[0] # Title Slide
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = "Unit Tesztelés a Gyakorlatban"
    subtitle.text = "Logika ellenőrzése Jest keretrendszerrel (Node.js)\n\n2025.05.xx.\nElőadó: [Te Neved]"

    # --- 2. DIA: Tesztelési Piramis ---
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]

    title.text = "Miért tesztelünk? A Tesztelési Piramis"
    content.text = (
        "• E2E (End-to-End):\n"
        "  - A teljes rendszer működése (pl. Cypress, Selenium).\n"
        "  - Lassú, drága.\n\n"
        "• Integration (Integrációs):\n"
        "  - Komponensek kapcsolata (pl. API + Adatbázis -> Postman).\n\n"
        "• Unit (Egységteszt):\n"
        "  - A legkisebb kódrészletek (függvények) izolált vizsgálata.\n"
        "  - Gyors, olcsó."
    )
    
    # Kép helyőrző szöveg hozzáadása
    txBox = slide.shapes.add_textbox(Inches(5.5), Inches(2), Inches(3), Inches(1))
    tf = txBox.text_frame
    tf.text = "[IDE ILLESSZ BE EGY KÉPET A TESZTELÉSI PIRAMISRÓL]"
    
    # Jegyzetek
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = "Oktatói megjegyzés: A vizsgafeladatban a Postman az integrációs teszt volt. Most lemegyünk a piramis aljára, a kód szintjére."

    # --- 3. DIA: Mi az a Unit Teszt? ---
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]

    title.text = "Mi az a Unit Teszt?"
    content.text = (
        "Definíció: Olyan módszer, ahol a forráskód egyedi egységeit (függvény) vizsgáljuk.\n\n"
        "Előnyök:\n"
        "• ⏱️ Gyors: Ezredmásodpercek alatt fut.\n"
        "• 🛡️ Biztonság: Bátran módosíthatod a kódot (Refactoring).\n"
        "• 📄 Dokumentáció: Megmutatja, hogyan kellene működnie a kódnak.\n\n"
        "Módszertan (AAA):\n"
        "1. Arrange (Előkészítés)\n"
        "2. Act (Végrehajtás)\n"
        "3. Assert (Ellenőrzés)"
    )

    # --- 4. DIA: A vizsgafeladat problémája ---
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]

    title.text = "A vizsgafeladat problémája (Scenario)"
    content.text = (
        "Feladat: A frontendnek ki kell írnia, hány éve aktív egy zenekar.\n\n"
        "Logika a dokumentációból:\n"
        "• Ha zárt intervallum (pl. \"1970-1991\"): Vég - Kezdet.\n"
        "• Ha nyitott (pl. \"1973-napjainkig\"): 2025 - Kezdet.\n\n"
        "Kérdés: Hogyan győződünk meg róla, hogy a számítás mindig helyes, anélkül, hogy mindig elindítanánk a szervert?\n\n"
        "Megoldás: Unit teszt!"
    )

    # --- 5. DIA: Az eszköz: Jest ---
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]

    title.text = "Az eszköz: Jest"
    content.text = (
        "Mi ez? A legnépszerűbb JavaScript tesztelő keretrendszer Node.js környezetben.\n\n"
        "Telepítés:\n"
        "npm install --save-dev jest\n\n"
        "Futtatás:\n"
        "npm test\n\n"
        "Fájl elnevezés:\n"
        "A Jest automatikusan keresi a *.test.js végű fájlokat."
    )

    # --- 6. DIA: 1. Lépés - A tiszta függvény ---
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]

    title.text = "1. Lépés - A tiszta függvény (utils.js)"
    
    # Kód beillesztése monospaced fonttal
    text_frame = content.text_frame
    text_frame.clear()
    p = text_frame.paragraphs[0]
    p.text = "Hozzunk létre egy utils.js fájlt (adatbázis nélkül!):"
    
    p = text_frame.add_paragraph()
    p.text = (
        "function szamolAktivEveket(aktivEvekStr) {\n"
        "    const darabok = aktivEvekStr.split('-');\n"
        "    const kezdoEv = parseInt(darabok[0]);\n"
        "    let vegEv;\n\n"
        "    if (darabok[1].includes(\"napjainkig\")) {\n"
        "        vegEv = 2025; // Vizsga szerinti fix év\n"
        "    } else {\n"
        "        vegEv = parseInt(darabok[1]);\n"
        "    }\n"
        "    return vegEv - kezdoEv;\n"
        "}\n"
        "module.exports = szamolAktivEveket;"
    )
    p.font.name = "Courier New"
    p.font.size = Pt(14)

    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = "Oktatói megjegyzés: Emeld ki a split és includes függvényeket."

    # --- 7. DIA: 2. Lépés - A Tesztfájl ---
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]

    title.text = "2. Lépés - A Tesztfájl (utils.test.js)"
    
    text_frame = content.text_frame
    text_frame.clear()
    p = text_frame.paragraphs[0]
    p.text = "A teszt szerkezete Jest-ben:"
    
    p = text_frame.add_paragraph()
    p.text = (
        "const szamolAktivEveket = require('./utils');\n\n"
        "describe('Aktív évek számítása', () => {\n\n"
        "    test('Lezárt időszak (pl. Queen)', () => {\n"
        "        // Arrange & Act & Assert egyben:\n"
        "        expect(szamolAktivEveket(\"1970-1991\")).toBe(21);\n"
        "    });\n\n"
        "});"
    )
    p.font.name = "Courier New"
    p.font.size = Pt(14)
    
    p = text_frame.add_paragraph()
    p.text = "\nMagyarázat:\n• describe: Tesztek csoportosítása\n• test: Egy konkrét eset\n• expect(...).toBe(...): Az elvárt eredmény vizsgálata"
    p.font.size = Pt(18)

    # --- 8. DIA: További tesztesetek ---
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]

    title.text = "További tesztesetek (Corner cases)"
    
    text_frame = content.text_frame
    text_frame.clear()
    p = text_frame.paragraphs[0]
    p.text = "Mi van, ha a zenekar még aktív?"
    
    p = text_frame.add_paragraph()
    p.text = (
        "test('Aktív zenekar (napjainkig -> 2025)', () => {\n"
        "    const bemenet = \"1973-napjainkig\";\n"
        "    const elvart = 52; // 2025 - 1973\n\n"
        "    const eredmeny = szamolAktivEveket(bemenet);\n\n"
        "    expect(eredmeny).toBe(elvart);\n"
        "});"
    )
    p.font.name = "Courier New"
    p.font.size = Pt(14)

    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = "Oktatói megjegyzés: Ez a teszt garantálja, hogy a 'napjainkig' szó kezelése működik."

    # --- 9. DIA: Eredmény értelmezése ---
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]

    title.text = "Eredmény értelmezése"
    
    text_frame = content.text_frame
    text_frame.clear()
    p = text_frame.paragraphs[0]
    p.text = "Kimenet a terminálban:"
    
    p = text_frame.add_paragraph()
    p.text = (
        "PASS  ./utils.test.js\n"
        "Aktív évek számítása\n"
        "  √ Lezárt időszak (pl. Queen) (2 ms)\n"
        "  √ Aktív zenekar (napjainkig -> 2025) (1 ms)\n\n"
        "Test Suites: 1 passed, 1 total\n"
        "Tests:       2 passed, 2 total"
    )
    p.font.name = "Courier New"
    p.font.size = Pt(16)
    p.font.bold = True
    
    p = text_frame.add_paragraph()
    p.text = "\nHa hiba van (PIROS): A Jest megmondja, mit várt (Expected) és mit kapott (Received)."
    p.font.size = Pt(18)
    
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = "Hangsúlyozd: Ha a jövőben a vizsgaszabályzat 2026-ra változik, csak átírod a kódot, lefuttatod a tesztet, és azonnal látod, ha elfelejtetted valahol frissíteni."

    # --- 10. DIA: Összefoglalás ---
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]

    title.text = "Összefoglalás"
    content.text = (
        "Mit tanultunk?\n"
        "• A Unit teszt a kód logikáját vizsgálja izoláltan.\n"
        "• A Postman (előző óra) a végpontokat és az adatbázist vizsgálja.\n"
        "• A Jest segítségével automatizálhatjuk az ellenőrzést.\n\n"
        "Következő lépés:\n"
        "Készítsétek el a utils.test.js fájlt a projekt mappájában és futtassátok le!"
    )

    # Mentés
    prs.save('Unit_Teszteles_Jest.pptx')
    print("A prezentáció elkészült: Unit_Teszteles_Jest.pptx")

if __name__ == "__main__":
    create_presentation()